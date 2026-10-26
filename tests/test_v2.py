"""V2 regression uses synthetic data only. Automatic approval below is never for real cases."""
import io,json,zipfile,collections,re,threading,urllib.request,urllib.error,http.cookiejar,time
from pathlib import Path
import pytest
from cloak.scenarios import scenario_from,suggest_purpose
from cloak.organize import propose,set_assignments
from cloak.core import Detector,TOKEN,Hit
from cloak.formats import Document,_ocr
from cloak.workflow import scan,review_update,redact,restore,verify,SOFT_ISSUES,_attach,input_data
from cloak.vault import open_vault,seal
from cloak.teaching import student_copy
from cloak.exporter import markdown
from fixtures import TEXT,DICT,PHONE,EMAIL,PERSON,docx,pptx,image
PW='SYNTHETIC-v2-password-2026'

@pytest.mark.parametrize('purpose,md,dates',[('ai',True,True),('students',False,True),('public',False,False),('classroom',False,True),('external',False,True)])
def test_scenarios(purpose,md,dates):
    s=scenario_from({'purpose':purpose});assert s.export_md==md and s.preserve_dates==dates
@pytest.mark.parametrize('data',[{'role':'unknown'},{'ppt_pdf':'yes'},{'organization':'guess'},{'unexpected':1}])
def test_invalid_scenarios(data):
    with pytest.raises(ValueError):scenario_from(data)
def test_purpose_suggestion():assert suggest_purpose('课件要发给学生')=='students'

def record(text,sha='x',suffix='.txt'):
    return {'units':[{'text':text,'kind':'text'}],'sha256':sha,'suffix':suffix}
def test_conservative_groups():
    files=[record('案号：（2026）京0105民初123号','1'),record('合同。案号：（2026）京0105民初123号','2'),record('原告：赵测试','3'),record('案号：（2026）京0105民初123号及（2026）沪0105民初456号','4'),record('法律法规','5')]
    result=propose(files);assert files[0]['case_id']==files[1]['case_id']=='CASE001';assert files[2]['case_id']==files[3]['case_id']=='PENDING';assert files[4]['case_id']=='COMMON';assert result['pending']==2
def test_duplicates_retained():
    files=[record('合同编号：AB-001'),record('合同编号：AB-001')];propose(files);assert len(files)==2 and files[1]['duplicate_of']==0

def approve(ws,assignments=None):
    job=open_vault(ws/'private/task.lpcv',PW)
    if job.get('scenario'):set_assignments(job,assignments or [{'index':i,'case_id':'CASE001','category':f['category']} for i,f in enumerate(job['files'])])
    review_update(job,confirmations=list(SOFT_ISSUES));seal(ws/'private/task.lpcv',job,PW);return job

def test_batch_ner_equal():
    texts=TEXT.splitlines();a=Detector(dictionary=DICT);b=Detector(dictionary=DICT);b.prime(texts)
    for t in texts:assert [h.dict() for h in a.find(t)]==[h.dict() for h in b.find(t)]

def test_sparse_review_and_md(tmp_path):
    p=tmp_path/'秘密原文件名.docx';docx(p);ws=tmp_path/'job';scan([p],ws,PW,dictionary=DICT,scenario={'purpose':'ai'})
    j=open_vault(ws/'private/task.lpcv',PW);d=Document(p.read_bytes(),'.docx');assert len(j['files'][0]['units'])<len(d.units)/5;_attach(d,j['files'][0]);approve(ws)
    result=redact(ws,PW);assert result['status']=='CHECKED'
    md=next((ws/'share').rglob('file-001.docx.md')).read_text();assert PERSON not in md and PHONE not in md and '秘密原文件名' not in md and '来源' in md and '修订' in md
    assert '2026年10月23日' in (ws/'share/日期索引.md').read_text();assert verify(ws,PW)['status']=='CHECKED'
    restored=tmp_path/'restored.docx';restore(ws,PW,restored);assert restored.read_bytes()==p.read_bytes()

def test_group_gate_and_identity_scope(tmp_path):
    paths=[]
    for i in range(3):p=tmp_path/f'{i}.txt';p.write_text('原告：赵测试，电话：'+PHONE);paths.append(p)
    ws=tmp_path/'job';scan(paths,ws,PW,scenario={'purpose':'ai'},dictionary={})
    j=open_vault(ws/'private/task.lpcv',PW)
    with pytest.raises(ValueError,match='GROUP_REVIEW_REQUIRED'):review_update(j,confirmations=list(SOFT_ISSUES))
    approve(ws,[{'index':i,'case_id':'CASE001' if i<2 else 'CASE002'} for i in range(3)])
    assert redact(ws,PW)['status']=='CHECKED'
    texts=[p.read_text() for p in sorted((ws/'share').rglob('*.txt'))]
    assert TOKEN.findall(texts[0])==TOKEN.findall(texts[1]) and TOKEN.findall(texts[0])!=TOKEN.findall(texts[2])

def test_same_name_manual_entities(tmp_path):
    p=tmp_path/'s.txt';p.write_text('原告：赵测试，被告：赵测试。');ws=tmp_path/'job';scan([p],ws,PW,scenario={'organization':'single'})
    j=open_vault(ws/'private/task.lpcv',PW);h=j['files'][0]['units'][0]['hits'];assert len(h)==2
    h[0]['entity']='person-a';h[1]['entity']='person-b';set_assignments(j,[{'index':0,'case_id':'CASE001'}]);review_update(j,confirmations=list(SOFT_ISSUES));seal(ws/'private/task.lpcv',j,PW);redact(ws,PW)
    tokens=TOKEN.findall(next((ws/'share').rglob('*.txt')).read_text());assert len(set(tokens))==2

def test_manual_image_never_leaks_to_md():
    d=Document(image(),'.png');u=next(u for u in d.units if PHONE in u.text);box=u.boxes[0];d.manual_rectangles=[{'image':'image','box':[0,box[3]-5,1600,box[5]+5]}]
    result=markdown(d,lambda u,h:'MASK',0).decode();assert PHONE not in result and '[手工遮挡区域' in result

def test_student_purge_and_roundtrip(tmp_path):
    p=tmp_path/'teaching.pptx';pptx(p,True);original=p.read_bytes();clean,counts=student_copy(original)
    from pptx import Presentation
    assert len(Presentation(io.BytesIO(clean)).slides)==1 and counts['hidden_slides_removed']==1
    with zipfile.ZipFile(io.BytesIO(clean)) as z:
        assert not any('notesSlide' in n or 'notesMaster' in n for n in z.namelist())
        assert not any(b'\xe9\x9a\x90\xe8\x97\x8f\xe9\xa1\xb5' in z.read(n) for n in z.namelist() if n.endswith('.xml'))
    ws=tmp_path/'job';scan([p],ws,PW,dictionary=DICT,scenario={'purpose':'students','organization':'single'});approve(ws);r=redact(ws,PW);assert r['status']=='CHECKED'
    output=next((ws/'share').rglob('*.pptx'));assert len(Presentation(output).slides)==1
    out=tmp_path/'restored.pptx';restore(ws,PW,out);assert out.read_bytes()==original
    edited=tmp_path/'edited-restored.pptx';assert restore(ws,PW,edited,output)['status']=='RESTORED_EDITED';assert len(Presentation(edited).slides)==1

def test_derived_change_detected(tmp_path):
    p=tmp_path/'s.txt';p.write_text(PHONE);ws=tmp_path/'job';scan([p],ws,PW,scenario={'organization':'single'});approve(ws);redact(ws,PW)
    next((ws/'share').rglob('*.txt.md')).write_text('changed');assert verify(ws,PW)['status']=='NEEDS_REVIEW'

def test_public_preserve_options(tmp_path):
    p=tmp_path/'s.txt';p.write_text('2026年10月23日支付120000元。'+PHONE);ws=tmp_path/'job';scan([p],ws,PW,scenario={'purpose':'public','preserve_dates':True,'preserve_amounts':True,'organization':'single'})
    approve(ws);assert redact(ws,PW)['status']=='CHECKED';text=next((ws/'share').rglob('*.txt')).read_text();assert '120000元' in text and '2026年10月23日' in text and PHONE not in text

def test_ocr_cache_independent_hits():
    a=Document(image(),'.png');a.units[0].hits=[Hit(0,1,'SECRET','test')];b=Document(image(),'.png');assert not b.units[0].hits and _ocr.cache_info().hits>0

def test_desktop_security_and_api(tmp_path,monkeypatch):
    from cloak.desktop import serve
    result={};event=threading.Event()
    def ready(url,server):result.update(url=url,server=server);event.set()
    t=threading.Thread(target=lambda:serve(open_browser=False,ready=ready,home=tmp_path/'tasks'),daemon=True);t.start();assert event.wait(15)
    base=result['url'].split('/?')[0];jar=http.cookiejar.CookieJar();client=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    try:
        with pytest.raises(urllib.error.HTTPError):urllib.request.urlopen(base+'/state')
        client.open(result['url']).read();s=json.load(client.open(base+'/state'));assert 'password' not in s and 'job' not in s
        def post(path,data,csrf=s['csrf'],origin=base):return client.open(urllib.request.Request(base+path,data=json.dumps(data).encode(),headers={'Content-Type':'application/json','Origin':origin,'X-CSRF':csrf}))
        with pytest.raises(urllib.error.HTTPError):post('/reset',{},csrf='wrong')
        with pytest.raises(urllib.error.HTTPError):post('/reset',{},origin='https://evil.test')
        with pytest.raises(urllib.error.HTTPError):client.open(base+'/../scripts/cloak/vault.py')
        p=tmp_path/'synthetic.txt';p.write_text(PHONE);post('/paths',{'paths':[str(p)]}).read();post('/scan',{'scenario':{'organization':'single'},'password':PW,'password_again':PW}).read()
        for _ in range(120):
            s=json.load(client.open(base+'/state'))
            if s['status']!='BUSY':break
            time.sleep(.2)
        assert s['status']=='REVIEW';assert 'password' not in json.dumps(s)
        post('/groups',{'assignments':[{'index':0,'case_id':'CASE001'}]}).read()
        with pytest.raises(urllib.error.HTTPError):post('/export',{'confirmations':list(SOFT_ISSUES)})
        post('/seen',{'fi':0,'confirmed':True}).read();post('/export',{'confirmations':list(SOFT_ISSUES)}).read()
        for _ in range(120):
            s=json.load(client.open(base+'/state'))
            if s['status']!='BUSY':break
            time.sleep(.2)
        assert s['status']=='DONE' and s['result']['verification']=='CHECKED'
        import subprocess
        calls=[]
        with monkeypatch.context() as patch:
            patch.setattr('cloak.desktop.open_folder',lambda path:calls.append(str(path)))
            post('/open',{'kind':'share'}).read()
        assert calls==[str(Path(s['workspace'])/'share')]
        assert s['delivery_path']==str(Path(s['workspace'])/'share')
        restored=json.load(post('/restore',{'fi':0}));assert Path(restored['output']).read_text()==PHONE
    finally:result['server'].shutdown();t.join(15)

def test_recipient_public_enforces_defaults():
    s=scenario_from({'purpose':'ai','recipient':'public'});assert s.policy=='public' and not s.preserve_dates and not s.preserve_amounts
    assert scenario_from({'purpose':'external','recipient':'students'}).clean_student_copy

def test_group_transaction_atomic():
    j={'files':[{'case_id':'PENDING','category':'other'},{'case_id':'PENDING','category':'other'}]}
    import copy
    before=copy.deepcopy(j)
    with pytest.raises(ValueError):set_assignments(j,[{'index':0,'case_id':'CASE001'},{'index':1,'case_id':'secret real name'}])
    assert j==before

def test_md_second_pass_detects_unmarked_source():
    from cloak.exporter import verify_markdown
    d=Document('姓名：赵测试，电话：13812345678。```'.encode(),'.txt')
    md=markdown(d,lambda u,h:'MASK',0)
    v=verify_markdown(md,Detector(dictionary=DICT),{},[]);assert v['status']=='NEEDS_REVIEW' and v['residual_count']>0

def test_derived_pending_cannot_become_checked(tmp_path):
    p=tmp_path/'s.txt';p.write_text(PHONE);ws=tmp_path/'job';scan([p],ws,PW,scenario={'organization':'single'});approve(ws);redact(ws,PW)
    j=open_vault(ws/'private/task.lpcv',PW);j['verification'].append({'derived':'test','status':'NEEDS_REVIEW','residual_count':1,'residuals':[{'unit':'test','kind':'SECRET','rule':'test'}],'blockers':[]});seal(ws/'private/task.lpcv',j,PW)
    assert verify(ws,PW)['status']=='NEEDS_REVIEW'

@pytest.mark.parametrize('repetition',range(3))
def test_office_image_repeat(tmp_path,repetition):
    p=tmp_path/'s.docx';docx(p,True);ws=tmp_path/'job';scan([p],ws,PW,dictionary=DICT,scenario={'organization':'single'});approve(ws);r=redact(ws,PW);assert r['status']=='CHECKED',r['verification']

def test_desktop_reactivation_single_use(tmp_path):
    from cloak.desktop import serve,reactivate
    record={};event=threading.Event();descriptor=tmp_path/'session.json'
    def ready(url,server):record.update(url=url,server=server);event.set()
    thread=threading.Thread(target=lambda:serve(open_browser=False,ready=ready,home=tmp_path/'tasks',singleton_file=descriptor),daemon=True);thread.start();assert event.wait(10)
    try:
        initial=json.loads(descriptor.read_text());url=reactivate(descriptor);assert url and url!=record['url'] and str(initial['port']) in url
        with pytest.raises(urllib.error.HTTPError):urllib.request.urlopen(record['url'])
        jar=http.cookiejar.CookieJar();client=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar));client.open(url).read();assert json.load(client.open(url.split('/?')[0]+'/state'))['status']=='IDLE'
        with pytest.raises(urllib.error.HTTPError):urllib.request.urlopen(url)
        # Repeated activation returns the same existing engine port, never an external URL.
        again=reactivate(descriptor);assert again and again!=url and str(initial['port']) in again
    finally:record['server'].shutdown();thread.join(10)
    assert not descriptor.exists()

def test_exported_without_candidates_still_frozen():
    from cloak.desktop import exported
    assert exported({'mapping':{},'files':[{'output_name':'file-001.txt'}]})
    assert not exported({'mapping':{},'files':[{}]})
