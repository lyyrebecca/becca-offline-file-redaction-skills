import collections,io,json,os,re,socket,sys,threading,urllib.request,urllib.parse,http.cookiejar,zipfile
from pathlib import Path
import pytest
from lxml import etree as E
from cloak.core import Detector,Hit,Tokenizer,TOKEN,id_checksum,luhn,apply_spans
from cloak.vault import seal,open_vault,unb64
from cloak.formats import Document
from cloak.workflow import scan,review_update,redact,verify,restore,add_manual,SOFT_ISSUES
from fixtures import PERSON,PHONE,EMAIL,SECRET,ID,DICT,TEXT,make_all,docx,xlsx,pptx,image
PW='SYNTHETIC-test-only-2026'

@pytest.fixture(scope='session')
def detector():return Detector(dictionary=DICT)

@pytest.fixture
def autojob(tmp_path):
    def run(path,dictionary=DICT,policy='ai'):
        ws=tmp_path/('job'+str(len(list(tmp_path.glob('job*')))))
        scan([path],ws,PW,policy,dictionary)
        job=open_vault(ws/'private/task.lpcv',PW)
        confirmations=sorted(set(i for f in job['files'] for i in f['issues'])&SOFT_ISSUES)
        review_update(job,confirmations=confirmations)
        seal(ws/'private/task.lpcv',job,PW)
        result=redact(ws,PW)
        return ws,result
    return run


def test_model_local(detector):assert detector.ner_available
@pytest.mark.parametrize('kind,value',[('PHONE',PHONE),('EMAIL',EMAIL),('PERSON',PERSON),('SECRET',SECRET),('ID',ID)])
def test_known_entities(detector,kind,value):assert any(h.kind==kind and TEXT[h.start:h.end]==value for h in detector.find(TEXT))

def test_invalid_id_not_ignored(detector):assert any('invalid_checksum' in h.rule for h in detector.find(ID))
def test_id_checksum():assert id_checksum('11010519491231002X')
def test_luhn():assert luhn('4111111111111111') and not luhn('4111111111111112')
def test_width_normalization(detector):assert any(h.kind=='PHONE' for h in detector.find('１３８１２３４５６７８'))
def test_placeholder_protection(detector):assert not detector.find('[P:123456789abc:PERSON:001]')
def test_policy_public():assert {'DATE','AMOUNT'}.issubset({h.kind for h in Detector('public',use_ner=False).find('2026年10月23日支付120000元')})
def test_allow_high_conflict():
    hs=Detector(dictionary={'allow':[PHONE]},use_ner=False).find(PHONE);assert hs[0].mask and hs[0].reason=='high_risk_allow_conflict'
def test_token_scope():
    t=Tokenizer('123456789abc');a=t.token(PERSON,'PERSON','',0)
    assert t.token(PERSON,'PERSON','',0)==a and t.token(PERSON,'PERSON','',1)!=a
    assert t.token(PERSON,'PERSON','x',0)==t.token(PERSON,'PERSON','x',1)
def test_exact_alias_restoration():
    t=Tokenizer('123456789abc');a=t.token('赵某','PERSON','x',0);b=t.token('赵先生','PERSON','x',0)
    assert a!=b and t.mapping[a]=='赵某' and t.mapping[b]=='赵先生'
def test_span_reject():
    with pytest.raises(ValueError):apply_spans('abc',[Hit(1,4,'SECRET','x')],lambda h:'X')

@pytest.mark.parametrize('mutation',['wrong_password','tamper','truncated'])
def test_vault_authentication(tmp_path,mutation):
    p=tmp_path/'x';seal(p,{'secret':PERSON},PW);assert PERSON.encode() not in p.read_bytes()
    if mutation=='tamper':p.write_bytes(p.read_bytes()[:-1]+bytes([p.read_bytes()[-1]^1]))
    if mutation=='truncated':p.write_bytes(p.read_bytes()[:15])
    with pytest.raises(Exception):open_vault(p,'different-passphrase' if mutation=='wrong_password' else PW)
def test_permissions(tmp_path):
    p=tmp_path/'x';seal(p,{},PW)
    if __import__('os').name!='nt':assert p.stat().st_mode&0o777==0o600
    assert open_vault(p,PW)=={}

@pytest.mark.parametrize('suffix',['.txt','.md','.docx','.xlsx','.pptx'])
def test_text_office_roundtrip(tmp_path,autojob,suffix):
    p=tmp_path/('source'+suffix)
    if suffix in ('.txt','.md'):p.write_text(TEXT)
    else:{'.docx':docx,'.xlsx':xlsx,'.pptx':pptx}[suffix](p)
    ws,report=autojob(p)
    out=ws/'share'/('file-001'+suffix);d=Document(out.read_bytes(),suffix);text='\n'.join(u.text for u in d.units)
    assert PHONE not in text and PERSON not in text and SECRET not in text
    assert '120000' in text and report['status']=='CHECKED',report
    assert verify(ws,PW)['status']=='CHECKED'
    restored=tmp_path/('restored'+suffix);r=restore(ws,PW,restored,edited=out)
    assert r['status']=='RESTORED_EDITED'
    got='\n'.join(u.text for u in Document(restored.read_bytes(),suffix).units)
    assert PERSON in got and PHONE in got
    archive=tmp_path/('archive'+suffix);restore(ws,PW,archive);assert archive.read_bytes()==p.read_bytes()

@pytest.mark.parametrize('suffix',['.png','.jpg','.pdf'])
def test_raster_roundtrip(tmp_path,autojob,suffix):
    from PIL import Image
    p=tmp_path/('source'+suffix);im=Image.open(io.BytesIO(image()));im.save(p,'PDF' if suffix=='.pdf' else 'PNG' if suffix=='.png' else 'JPEG')
    ws,report=autojob(p)
    d=Document((ws/'share'/('file-001'+suffix)).read_bytes(),suffix)
    text='\n'.join(u.text for u in d.units);assert PHONE not in text and EMAIL not in text
    assert report['status']=='CHECKED',report
    if suffix=='.pdf':
        from pypdf import PdfReader
        assert not ''.join(pg.extract_text() for pg in PdfReader(ws/'share/file-001.pdf').pages)
    out=tmp_path/('original'+suffix);restore(ws,PW,out);assert out.read_bytes()==p.read_bytes()


def test_scan_has_no_plaintext(tmp_path):
    p=tmp_path/'private_name.txt';p.write_text(TEXT);ws=tmp_path/'job';result=scan([p],ws,PW,dictionary=DICT)
    assert result['status']=='NEEDS_REVIEW'
    for f in ws.rglob('*'):
        if f.is_file():
            raw=f.read_bytes();assert PHONE.encode() not in raw and PERSON.encode() not in raw and str(p).encode() not in raw
    with pytest.raises(ValueError,match='LOCAL_REVIEW_REQUIRED'):redact(ws,PW)

def test_keep_reason_required(tmp_path):
    p=tmp_path/'a.txt';p.write_text(PHONE);ws=tmp_path/'job';scan([p],ws,PW);job=open_vault(ws/'private/task.lpcv',PW)
    with pytest.raises(ValueError,match='KEEP_REQUIRES_REASON'):review_update(job,{'0:0:0':{'mask':False}},['ENTITY_AMBIGUITY_REVIEW'])
def test_manual_add(tmp_path):
    p=tmp_path/'a.txt';p.write_text('XYZ-unknown');ws=tmp_path/'job';scan([p],ws,PW);j=open_vault(ws/'private/task.lpcv',PW)
    add_manual(j,0,0,0,3,'SECRET');assert j['files'][0]['units'][0]['hits'][0]['rule']=='local_manual'
    with pytest.raises(ValueError,match='SPAN_OVERLAP'):add_manual(j,0,0,0,2,'SECRET')
def test_keep_authorization(tmp_path):
    p=tmp_path/'a.txt';p.write_text(PHONE);ws=tmp_path/'job';scan([p],ws,PW);j=open_vault(ws/'private/task.lpcv',PW)
    review_update(j,{'0:0:0':{'mask':False,'reason':'合成测试经许可保留'}},['ENTITY_AMBIGUITY_REVIEW']);seal(ws/'private/task.lpcv',j,PW)
    r=redact(ws,PW);assert r['status']=='CHECKED' and r['kept_candidates']==1

def test_restoration_preserves_edits(tmp_path,autojob):
    p=tmp_path/'s.docx';docx(p);ws,_=autojob(p)
    with zipfile.ZipFile(ws/'share/file-001.docx') as z:parts={n:z.read(n) for n in z.namelist()}
    root=E.fromstring(parts['word/document.xml']);n=root.xpath('//*[local-name()="t"]')[-1];n.text=(n.text or '')+' 合成审阅意见：建议补充违约责任。';parts['word/document.xml']=E.tostring(root)
    edit=tmp_path/'edited.docx'
    with zipfile.ZipFile(edit,'w') as z:
        for n,b in parts.items():z.writestr(n,b)
    out=tmp_path/'restored.docx';r=restore(ws,PW,out,edit);assert r['status']=='RESTORED_EDITED'
    text=''.join(u.text for u in Document(out.read_bytes(),'.docx').units)
    assert '建议补充违约责任' in text and PERSON in text
    with zipfile.ZipFile(out) as z:assert 'word/comments.xml' in z.namelist() and b'delText' in z.read('word/document.xml')

@pytest.mark.parametrize('change',['delete','copy','damage','foreign'])
def test_restore_conflicts(tmp_path,autojob,change):
    p=tmp_path/'s.txt';p.write_text('原告：'+PERSON+'，电话：'+PHONE);ws,_=autojob(p)
    text=(ws/'share/file-001.txt').read_text();m=TOKEN.search(text);assert m
    if change=='delete':text=text[:m.start()]+text[m.end():]
    elif change=='copy':text+=' '+m.group()
    elif change=='damage':text=text.replace(m.group(),m.group().replace('PERSON','PERS0N'))
    else:text+=' [P:ffffffffffff:PERSON:001]'
    edit=tmp_path/'edit.txt';edit.write_text(text);out=tmp_path/'restore.txt';r=restore(ws,PW,out,edit)
    assert r['status']=='NEEDS_REVIEW' and not out.exists()

def test_changed_output_verification(tmp_path,autojob):
    p=tmp_path/'s.txt';p.write_text(PHONE);ws,_=autojob(p)
    f=ws/'share/file-001.txt';f.write_text(f.read_text()+' '+PHONE)
    assert verify(ws,PW)['status']=='NEEDS_REVIEW'
def test_source_unchanged(tmp_path,autojob):
    p=tmp_path/'s.txt';p.write_text(TEXT);original=p.read_bytes();ws,_=autojob(p);assert p.read_bytes()==original
    with pytest.raises(ValueError,match='OUTPUT_ALREADY_EXISTS'):restore(ws,PW,p)

def modify_zip(path,name,data):
    with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    parts[name]=data
    with zipfile.ZipFile(path,'w') as z:
        for n,b in parts.items():z.writestr(n,b)

@pytest.mark.parametrize('part',['word/embeddings/oleObject1.bin','customXml/item1.xml','word/charts/chart1.xml','_xmlsignatures/sig1.xml'])
def test_unsupported_parts_block(tmp_path,part):
    p=tmp_path/'s.docx';docx(p);modify_zip(p,part,b'<x>synthetic</x>' if part.endswith('.xml') else b'opaque')
    ws=tmp_path/'job';scan([p],ws,PW,dictionary=DICT);j=open_vault(ws/'private/task.lpcv',PW)
    with pytest.raises(ValueError,match='UNRESOLVED_STRUCTURAL_BLOCKERS'):review_update(j,confirmations=['ENTITY_AMBIGUITY_REVIEW'])
def test_sensitive_formula_blocked(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'s.xlsx';w=Workbook();w.active['A1']='="'+PHONE+'"';w.save(p)
    ws=tmp_path/'job';scan([p],ws,PW);j=open_vault(ws/'private/task.lpcv',PW)
    assert 'SENSITIVE_FORMULA_BLOCKED' in j['files'][0]['issues']
def test_xml_dtd_blocked(tmp_path):
    p=tmp_path/'s.docx';docx(p);modify_zip(p,'word/evil.xml',b'<!DOCTYPE x [<!ENTITY xx SYSTEM "file:///etc/passwd">]><x>&xx;</x>')
    with pytest.raises(ValueError,match='XML_DTD_FORBIDDEN'):Document(p.read_bytes(),'.docx')
def test_metadata_cleaned(tmp_path,autojob):
    p=tmp_path/'s.docx';docx(p);ws,_=autojob(p)
    with zipfile.ZipFile(ws/'share/file-001.docx') as z:assert PERSON.encode() not in z.read('docProps/core.xml') and SECRET.encode() not in z.read('docProps/core.xml')
def test_no_network_needed(tmp_path,monkeypatch,autojob):
    def deny(*a,**kw):raise AssertionError('NETWORK_FORBIDDEN')
    monkeypatch.setattr(socket.socket,'connect',deny);monkeypatch.setattr(socket,'create_connection',deny)
    p=tmp_path/'s.docx';docx(p);ws,r=autojob(p);assert r['status']=='CHECKED'

def test_review_security_and_approval(tmp_path):
    from cloak.review import serve
    p=tmp_path/'s.txt';p.write_text(PHONE);ws=tmp_path/'job';scan([p],ws,PW)
    ready=threading.Event();state={}
    def callback(url,server):state.update(url=url,server=server);ready.set()
    thread=threading.Thread(target=serve,args=(ws,PW),kwargs={'open_browser':False,'ready':callback},daemon=True);thread.start();assert ready.wait(30)
    url=state['url'];base=url.split('/?')[0];cj=http.cookiejar.CookieJar();opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    try:
        with pytest.raises(urllib.error.HTTPError):urllib.request.urlopen(base)
        page=opener.open(url).read().decode();assert PHONE in page and 'no-store' in opener.open(base).headers['Cache-Control']
        csrf=re.search(r'name="csrf" value="([^"]+)"',page).group(1)
        with pytest.raises(urllib.error.HTTPError):opener.open(base+'/approve',data=b'csrf=wrong')
        for origin,site in [('https://evil.example','cross-site'),('null','cross-site')]:
            req=urllib.request.Request(base+'/approve',data=urllib.parse.urlencode({'csrf':csrf}).encode(),headers={'Origin':origin,'Sec-Fetch-Site':site})
            with pytest.raises(urllib.error.HTTPError):opener.open(req)
        page=opener.open(base+'/decision',data=urllib.parse.urlencode({'csrf':csrf,'pos':0,'mask0':'on','reason0':''}).encode()).read()
        page=opener.open(base+'/approve',data=urllib.parse.urlencode({'csrf':csrf,'confirm':'ENTITY_AMBIGUITY_REVIEW'}).encode()).read()
        thread.join(10);assert open_vault(ws/'private/task.lpcv',PW)['reviewed']
    finally:state['server'].shutdown()


def test_review_missing_ner_does_not_downgrade(tmp_path,monkeypatch):
    import cloak.workflow as wf
    class Missing(Detector):
        def __init__(self,*a,**kw):super().__init__(*a,use_ner=False,**kw)
    monkeypatch.setattr(wf,'Detector',Missing)
    p=tmp_path/'s.txt';p.write_text(PHONE)
    with pytest.raises(ValueError,match='LOCAL_NER_UNAVAILABLE'):wf.scan([p],tmp_path/'job',PW)

def test_all_office_raster_media(tmp_path,autojob):
    p=tmp_path/'s.docx';docx(p,with_image=True);ws,r=autojob(p)
    assert r['status']=='CHECKED'
    out=tmp_path/'restore.docx';rr=restore(ws,PW,out,ws/'share/file-001.docx');assert rr['status']=='RESTORED_EDITED'
    with zipfile.ZipFile(p) as a,zipfile.ZipFile(out) as b:
        for name in a.namelist():
            if name.startswith('word/media/'):assert a.read(name)==b.read(name)

def test_formula_preserved(tmp_path,autojob):
    from openpyxl import load_workbook
    p=tmp_path/'s.xlsx';xlsx(p);ws,_=autojob(p)
    w=load_workbook(ws/'share/file-001.xlsx');assert w.active['D2'].value=='=C2*2' and w.active['C2'].value==120000
    assert w['内部'].sheet_state=='hidden'

def test_metadata_independent_verifier(tmp_path,autojob):
    p=tmp_path/'s.docx';docx(p);ws,_=autojob(p)
    out=ws/'share/file-001.docx'
    with zipfile.ZipFile(out) as z:root=E.fromstring(z.read('docProps/core.xml'))
    root.xpath('//*[local-name()="creator"]')[0].text=PERSON
    modify_zip(out,'docProps/core.xml',E.tostring(root))
    r=verify(ws,PW);assert r['status']=='NEEDS_REVIEW' and any(x['kind']=='METADATA' for x in r['files'][0]['residuals'])

def test_sensitive_sheet_reference_blocked(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'s.xlsx';w=Workbook();w.active.title=PERSON;w.save(p)
    ws=tmp_path/'job';scan([p],ws,PW,dictionary=DICT);j=open_vault(ws/'private/task.lpcv',PW)
    assert 'SENSITIVE_REFERENCE_BLOCKED' in j['files'][0]['issues']

def test_manual_rectangle(tmp_path,autojob):
    from PIL import Image
    p=tmp_path/'s.png';p.write_bytes(image());ws=tmp_path/'job';scan([p],ws,PW,dictionary=DICT)
    j=open_vault(ws/'private/task.lpcv',PW);j['files'][0]['rectangles']=[{'image':'image','box':[0,0,100,100]}]
    review_update(j,confirmations=['IMAGE_VISUAL_REVIEW','ENTITY_AMBIGUITY_REVIEW']);seal(ws/'private/task.lpcv',j,PW)
    redact(ws,PW);im=Image.open(ws/'share/file-001.png');assert im.getpixel((50,50))==(0,0,0)

def test_legacy_macro_rejected(tmp_path):
    p=tmp_path/'s.docm';p.write_bytes(b'synthetic')
    with pytest.raises(ValueError,match='FORMAT_UNSUPPORTED'):scan([p],tmp_path/'job',PW)

def test_cli_no_tty_and_safe_error(tmp_path):
    import subprocess
    p=tmp_path/'s.txt';p.write_text(TEXT)
    cli=Path(__file__).resolve().parents[1]/'scripts/lpc.py'
    r=subprocess.run([sys.executable,str(cli),'scan',str(p),'--workspace',str(tmp_path/'job')],stdin=subprocess.DEVNULL,capture_output=True,text=True,timeout=15)
    assert r.returncode==1 and 'LOCAL_TTY_REQUIRED' in r.stdout
    assert PHONE not in r.stdout+r.stderr and PERSON not in r.stdout+r.stderr

def test_public_policy_semantics(detector):
    hs=detector.find('合同价款120000元，2026年10月23日履行。北京市朝阳区人民法院受理。')
    assert not any(h.kind in ('AMOUNT','DATE') for h in hs)

def test_batch_scope_default(tmp_path):
    a=tmp_path/'a.txt';b=tmp_path/'b.txt';a.write_text(PERSON);b.write_text(PERSON);ws=tmp_path/'job'
    d={'terms':[{'text':PERSON,'kind':'PERSON'}]};scan([a,b],ws,PW,dictionary=d);j=open_vault(ws/'private/task.lpcv',PW)
    review_update(j,confirmations=['ENTITY_AMBIGUITY_REVIEW','CROSS_FILE_IDENTITY_REVIEW']);seal(ws/'private/task.lpcv',j,PW);redact(ws,PW)
    assert (ws/'share/file-001.txt').read_text()!=(ws/'share/file-002.txt').read_text()


def test_ocr_edge_glyph_is_covered(tmp_path,autojob):
    from PIL import Image
    p=tmp_path/'s.png';p.write_bytes(image());ws,r=autojob(p)
    im=Image.open(ws/'share/file-001.png')
    # In this synthetic Arial fixture OCR's bbox excludes the leading '1' (x~242).
    # The actual pixel mask must cover that glyph, not merely make OCR stop reading the full number.
    assert im.getpixel((250,250))==(0,0,0)

def test_excel_modified_cells_wrap(tmp_path,autojob):
    from openpyxl import load_workbook
    p=tmp_path/'s.xlsx';xlsx(p);ws,r=autojob(p);w=load_workbook(ws/'share/file-001.xlsx')
    assert w.active['A2'].alignment.wrap_text and w.active.row_dimensions[2].height>=66


def test_numeric_identifier_becomes_valid_text_cell(tmp_path,autojob):
    from openpyxl import Workbook,load_workbook
    p=tmp_path/'s.xlsx';w=Workbook();w.active['A1']=int(PHONE);w.save(p)
    ws,r=autojob(p);out=load_workbook(ws/'share/file-001.xlsx');assert out.active['A1'].data_type=='s' and TOKEN.fullmatch(out.active['A1'].value)

def test_numeric_identifier_with_formulas_blocks(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'s.xlsx';w=Workbook();w.active['A1']=int(PHONE);w.active['B1']='=A1*2';w.save(p)
    ws=tmp_path/'job';scan([p],ws,PW);j=open_vault(ws/'private/task.lpcv',PW);assert 'SENSITIVE_NUMERIC_WITH_FORMULAS' in j['files'][0]['issues']


def test_sensitive_zip_member_identifier_blocked(tmp_path):
    p=tmp_path/'s.docx';docx(p);modify_zip(p,'word/'+PHONE+'.xml',b'<safe/>')
    ws=tmp_path/'job';scan([p],ws,PW);j=open_vault(ws/'private/task.lpcv',PW)
    assert 'SENSITIVE_REFERENCE_BLOCKED' in j['files'][0]['issues']

def test_sensitive_namespace_identifier_blocked(tmp_path):
    p=tmp_path/'s.docx';docx(p);modify_zip(p,'word/extra.xml',('<x xmlns="urn:'+PHONE+'"/>').encode())
    ws=tmp_path/'job';scan([p],ws,PW);j=open_vault(ws/'private/task.lpcv',PW)
    assert 'SENSITIVE_REFERENCE_BLOCKED' in j['files'][0]['issues']


def test_exif_orientation_normalized_and_stripped():
    from PIL import Image
    upright=Image.open(io.BytesIO(image()))
    raw=upright.transpose(Image.Transpose.ROTATE_90);exif=raw.getexif();exif[274]=6
    b=io.BytesIO();raw.save(b,'JPEG',exif=exif)
    doc=Document(b.getvalue(),'.jpg');clean=Image.open(io.BytesIO(doc.images['image']))
    assert clean.size==upright.size and not clean.getexif()
